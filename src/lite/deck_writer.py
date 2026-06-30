"""PPTX writer for PaperPresenter Lite deck specs."""

from __future__ import annotations

from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from src.lite.models import DeckSpec, SlideSpec, VisualAsset


class LitePPTXWriter:
    """Write a deterministic PPTX from a DeckSpec."""

    def write(self, deck: DeckSpec, output_path: str | Path) -> None:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        visual_map = deck.visual_map()
        for index, slide_spec in enumerate(deck.slides):
            self._add_slide(prs, slide_spec, visual_map, is_first=index == 0)

        prs.save(str(output_path))

    def _add_slide(
        self,
        prs: Presentation,
        spec: SlideSpec,
        visual_map: dict[str, VisualAsset],
        is_first: bool,
    ) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_title(slide, spec.title, is_first=is_first)

        visual = self._first_visual(spec, visual_map)
        if spec.layout == "title":
            self._add_bullets(slide, spec.bullets, Inches(1.0), Inches(2.2), Inches(11.3), Inches(3.5), Pt(24))
        elif visual and spec.layout == "visual_full":
            self._add_image(slide, visual, Inches(0.9), Inches(1.55), Inches(11.5), Inches(4.85))
            self._add_caption(slide, visual, Inches(1.0), Inches(6.35), Inches(11.3), Inches(0.45))
        elif visual:
            self._add_bullets(slide, spec.bullets, Inches(0.75), Inches(1.65), Inches(5.3), Inches(4.9), Pt(22))
            self._add_image(slide, visual, Inches(6.35), Inches(1.65), Inches(6.1), Inches(4.35))
            self._add_caption(slide, visual, Inches(6.2), Inches(6.1), Inches(6.35), Inches(0.6))
        else:
            self._add_bullets(slide, spec.bullets, Inches(0.9), Inches(1.75), Inches(11.6), Inches(5.1), Pt(24))

        if spec.notes or spec.speaker_intent:
            notes = slide.notes_slide.notes_text_frame
            notes.text = "\n".join(part for part in [spec.speaker_intent, spec.notes] if part)

    def _first_visual(self, spec: SlideSpec, visual_map: dict[str, VisualAsset]) -> Optional[VisualAsset]:
        for asset_id in spec.visual_asset_ids:
            if asset_id in visual_map:
                return visual_map[asset_id]
        return None

    def _add_title(self, slide, title: str, is_first: bool = False) -> None:
        box = slide.shapes.add_textbox(Inches(0.55), Inches(0.45), Inches(12.2), Inches(0.9))
        frame = box.text_frame
        frame.text = title
        paragraph = frame.paragraphs[0]
        paragraph.font.size = Pt(34 if is_first else 30)
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor(0x22, 0x2A, 0x35)

    def _add_bullets(self, slide, bullets, left, top, width, height, font_size) -> None:
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = 0

        for index, bullet in enumerate(bullets[:6]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = bullet
            paragraph.font.size = font_size
            paragraph.font.color.rgb = RGBColor(0x2D, 0x33, 0x3B)
            paragraph.space_after = Pt(10)

    def _add_image(self, slide, visual: VisualAsset, left, top, width, height) -> None:
        image_path = Path(visual.image_path)
        if not image_path.exists():
            return
        picture = slide.shapes.add_picture(str(image_path), left, top, width=width)
        if picture.height > height:
            scale = height / picture.height
            picture.width = int(picture.width * scale)
            picture.height = int(picture.height * scale)
        picture.left = int(left + (width - picture.width) / 2)
        picture.top = int(top + (height - picture.height) / 2)

    def _add_caption(self, slide, visual: VisualAsset, left, top, width, height) -> None:
        caption = f"{visual.caption} (source page {visual.source_page}, {visual.extraction_status})"
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.word_wrap = True
        frame.text = caption
        paragraph = frame.paragraphs[0]
        paragraph.font.size = Pt(11)
        paragraph.font.color.rgb = RGBColor(0x5C, 0x66, 0x70)
        paragraph.alignment = PP_ALIGN.CENTER
